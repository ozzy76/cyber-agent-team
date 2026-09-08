"""Augment scf.db with data only available in the SCF spreadsheet release.

The JSON release (preprocess_scf.py) covers 67 external frameworks. The
spreadsheet covers ~249 — including HIPAA, GDPR, SOC 2 (via AICPA TSC),
CCPA, CMMC 2.0, FedRAMP R5, and the rest of the missing frameworks.
The spreadsheet also carries:

  - The Authoritative Sources index: geography, publisher, full name,
    URL, FDI slug, and STRM URL for each framework.
  - SCR-CMM Level 0..5 maturity criteria per control — the OSCAL and
    JSON releases both omit these.

This script is a layer on top of preprocess_scf.py: run preprocess
first to build the base DB, then run this to add the spreadsheet-only
coverage. Idempotent — drops and rebuilds the tables it owns; does not
touch tables owned by preprocess_scf.py (controls, assessment_objectives,
control_evidence_refs, control_risk_threats, controls_fts).

Source : SCF spreadsheet release from https://tools.scfconnectlabs.com
         (e.g., "Secure Controls Framework (SCF) - 2026.1.1.xlsx").
Default DB     : ~/.local/share/bateam/grc/scf.db  (BATEAM_SCF_DB)
Default source : $BATEAM_SCF_XLSX (no in-repo default)

Usage:
    python augment_from_spreadsheet.py
    python augment_from_spreadsheet.py --source PATH
    python augment_from_spreadsheet.py --db PATH
    python augment_from_spreadsheet.py --stats

Schema additions:
    framework_metadata(slug, fdi, geography, publisher, full_name,
                      title, url, strm_url)
    maturity_criteria(control_id, cmm_level, criteria_text)

Schema replacements:
    control_frameworks — wiped and repopulated from spreadsheet. The
    spreadsheet is a strict superset of the JSON release for framework
    crosswalks (verified empirically).
"""

from __future__ import annotations

import argparse
import os
import pathlib
import re
import sqlite3
import sys
import time
from typing import Any

DEFAULT_DB = pathlib.Path("~/.local/share/bateam/grc/scf.db").expanduser()

# Match preprocess_scf.py exactly so the same framework name produces
# the same slug across both ingest paths.
_WS_RE = re.compile(r"\s+")
_SLUG_RE = re.compile(r"[^a-z0-9]+")


def norm_text(s: str) -> str:
    return _WS_RE.sub(" ", s.replace("\r\n", " ").replace("\r", " ").replace("\n", " ")).strip()


def slugify(s: str) -> str:
    s = norm_text(s).lower()
    return _SLUG_RE.sub("-", s).strip("-")


def split_multi(raw: Any) -> list[str]:
    """Spreadsheet cells pack multiple locators with literal \\n."""
    if raw is None:
        return []
    parts = re.split(r"[\r\n]+", str(raw))
    return [p.strip() for p in parts if p.strip()]


def resolve_db(arg: str | None) -> pathlib.Path:
    if arg:
        return pathlib.Path(arg).expanduser()
    return pathlib.Path(os.environ.get("BATEAM_SCF_DB", str(DEFAULT_DB))).expanduser()


def resolve_source(arg: str | None) -> pathlib.Path:
    if arg:
        return pathlib.Path(arg).expanduser()
    env = os.environ.get("BATEAM_SCF_XLSX")
    if env:
        return pathlib.Path(env).expanduser()
    sys.exit(
        "error: SCF spreadsheet not specified.\n"
        "  Pass --source PATH, or set BATEAM_SCF_XLSX.\n"
        "  Download from https://tools.scfconnectlabs.com — the file is\n"
        "  named 'Secure Controls Framework (SCF) - X.Y.Z.xlsx'."
    )


def ensure_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        DROP TABLE IF EXISTS framework_metadata;
        CREATE TABLE framework_metadata (
            slug       TEXT PRIMARY KEY,
            fdi        TEXT,
            geography  TEXT,
            publisher  TEXT,
            full_name  TEXT,
            title      TEXT,
            url        TEXT,
            strm_url   TEXT
        );

        DROP TABLE IF EXISTS maturity_criteria;
        CREATE TABLE maturity_criteria (
            control_id    TEXT NOT NULL REFERENCES controls(id),
            cmm_level     INTEGER NOT NULL,
            criteria_text TEXT NOT NULL,
            PRIMARY KEY (control_id, cmm_level)
        );
        CREATE INDEX IF NOT EXISTS idx_maturity_control ON maturity_criteria(control_id);

        -- Wipe and rebuild from spreadsheet (strict superset of JSON).
        DELETE FROM control_frameworks;
        """
    )


def load_authoritative_sources(wb) -> dict[str, dict[str, Any]]:
    """Return {normalized_header: {fdi, geography, publisher, full_name, title, url, strm_url}}.

    Column layout (header row 1):
      A Geography
      B SCF Column Header              ← join key into SCF 2026.1
      C Focal Document Identifier (FDI)
      D Source                          ← publisher
      E Focal Document Name (FDN)
      F Focal Document Title (FDT)
      G Focal Document Source (FDS)    ← public URL
      H Set Theory Relationship Mapping (STRM)
    """
    ws = wb["Authoritative Sources"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)  # noqa: F841 — assert layout if needed
    out: dict[str, dict[str, Any]] = {}
    for row in rows:
        if not row or not row[1]:
            continue
        n = norm_text(str(row[1]))
        out[n] = {
            "header": n,
            "fdi": (row[2] or "").strip() if row[2] else None,
            "geography": (row[0] or "").strip() if row[0] else None,
            "publisher": (row[3] or "").strip() if row[3] else None,
            "full_name": (row[4] or "").strip() if row[4] else None,
            "title": (row[5] or "").strip() if row[5] else None,
            "url": (row[6] or "").strip() if row[6] else None,
            "strm_url": (row[7] or "").strip() if row[7] else None,
        }
    return out


_CMM_LEVEL_RE = re.compile(r"^SCR-CMM Level (\d)\b", re.IGNORECASE)


def classify_headers(
    hdr_row: tuple,
    authoritative: dict[str, dict[str, Any]],
) -> tuple[int | None, list[tuple[int, str, str, dict[str, Any]]], list[tuple[int, int]]]:
    """Return (ctrl_id_col_idx, framework_cols, cmm_cols).

    framework_cols entries: (col_idx, slug, normalized_header, auth_meta)
    cmm_cols entries: (col_idx, level)
    """
    ctrl_id_col: int | None = None
    framework_cols: list[tuple[int, str, str, dict[str, Any]]] = []
    cmm_cols: list[tuple[int, int]] = []
    for i, h in enumerate(hdr_row):
        if h is None:
            continue
        n = norm_text(str(h))
        if n == "SCF #":
            ctrl_id_col = i
            continue
        m = _CMM_LEVEL_RE.match(n)
        if m:
            cmm_cols.append((i, int(m.group(1))))
            continue
        if n in authoritative:
            framework_cols.append((i, slugify(n), n, authoritative[n]))
    return ctrl_id_col, framework_cols, cmm_cols


def populate(
    conn: sqlite3.Connection,
    authoritative: dict[str, dict[str, Any]],
    framework_cols: list[tuple[int, str, str, dict[str, Any]]],
    cmm_cols: list[tuple[int, int]],
    ctrl_id_col: int,
    rows_iter,
) -> dict[str, int]:
    """Populate framework_metadata, maturity_criteria, and control_frameworks."""
    # Only insert metadata for the frameworks actually referenced by the
    # SCF 2026.1 header — keeps the metadata table aligned with the
    # control_frameworks slugs we actually emit.
    fm_rows = [
        (
            slug,
            meta.get("fdi"),
            meta.get("geography"),
            meta.get("publisher"),
            meta.get("full_name"),
            meta.get("title"),
            meta.get("url"),
            meta.get("strm_url"),
        )
        for (_col_idx, slug, _hdr, meta) in framework_cols
    ]
    # Dedupe by slug (same name → same slug for any duplicate column edge case)
    seen_slugs: set[str] = set()
    fm_unique = []
    for row in fm_rows:
        if row[0] in seen_slugs:
            continue
        seen_slugs.add(row[0])
        fm_unique.append(row)
    conn.executemany(
        "INSERT INTO framework_metadata"
        "(slug, fdi, geography, publisher, full_name, title, url, strm_url) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        fm_unique,
    )

    counts = {
        "framework_metadata": len(fm_unique),
        "control_frameworks": 0,
        "maturity_criteria": 0,
        "rows_walked": 0,
        "rows_with_id": 0,
    }
    cf_batch: list[tuple[str, str, str, str]] = []
    mc_batch: list[tuple[str, int, str]] = []
    BATCH = 5000

    def flush_cf():
        if cf_batch:
            conn.executemany(
                "INSERT INTO control_frameworks"
                "(control_id, framework_slug, framework_name, locator) "
                "VALUES (?, ?, ?, ?)",
                cf_batch,
            )
            counts["control_frameworks"] += len(cf_batch)
            cf_batch.clear()

    def flush_mc():
        if mc_batch:
            conn.executemany(
                "INSERT OR REPLACE INTO maturity_criteria"
                "(control_id, cmm_level, criteria_text) VALUES (?, ?, ?)",
                mc_batch,
            )
            counts["maturity_criteria"] += len(mc_batch)
            mc_batch.clear()

    # Map slug → display name (use the spreadsheet header, normalized)
    slug_to_name = {slug: hdr for (_i, slug, hdr, _m) in framework_cols}

    for row in rows_iter:
        counts["rows_walked"] += 1
        cid_raw = row[ctrl_id_col] if ctrl_id_col is not None and ctrl_id_col < len(row) else None
        if not cid_raw:
            continue
        cid = str(cid_raw).strip()
        if not cid:
            continue
        counts["rows_with_id"] += 1

        for col_idx, slug, _hdr, _meta in framework_cols:
            if col_idx >= len(row):
                continue
            val = row[col_idx]
            if val is None:
                continue
            for locator in split_multi(val):
                cf_batch.append((cid, slug, slug_to_name[slug], locator))
        if len(cf_batch) >= BATCH:
            flush_cf()

        for col_idx, level in cmm_cols:
            if col_idx >= len(row):
                continue
            criteria = row[col_idx]
            if criteria is None:
                continue
            text = str(criteria).strip()
            if text:
                mc_batch.append((cid, level, text))
        if len(mc_batch) >= BATCH:
            flush_mc()

    flush_cf()
    flush_mc()
    return counts


def print_stats(conn: sqlite3.Connection) -> None:
    rows = {
        "framework_metadata": conn.execute("SELECT COUNT(*) FROM framework_metadata").fetchone()[0],
        "control_frameworks": conn.execute("SELECT COUNT(*) FROM control_frameworks").fetchone()[0],
        "framework_slugs (distinct)": conn.execute(
            "SELECT COUNT(DISTINCT framework_slug) FROM control_frameworks"
        ).fetchone()[0],
        "maturity_criteria": conn.execute("SELECT COUNT(*) FROM maturity_criteria").fetchone()[0],
        "maturity_controls_covered": conn.execute(
            "SELECT COUNT(DISTINCT control_id) FROM maturity_criteria"
        ).fetchone()[0],
    }
    print("Augmentation row counts:")
    for k, v in rows.items():
        print(f"  {k:30s} {v:>10d}")

    # Spot-checks for previously-missing frameworks
    print("\nCoverage spot-check (controls per slug):")
    for slug in (
        "pci-dss-4-0-1",
        "aicpa-tsc-2017-2022-used-for-soc-2",
        "us-hipaa-administrative-simplification-2013",
        "us-hipaa-security-rule-nist-sp-800-66-r2",
        "emea-eu-gdpr",
        "us-ca-ccpa-2025",
        "us-cmmc-2-0-level-2",
        "us-fedramp-r5-moderate",
        "nist-csf-2-0",
        "iso-27001-2022",
    ):
        n = conn.execute(
            "SELECT COUNT(DISTINCT control_id) FROM control_frameworks WHERE framework_slug = ?",
            (slug,),
        ).fetchone()[0]
        present = "OK" if n > 0 else "MISS"
        print(f"  [{present:>4s}]  {slug:55s}  {n:>5d} controls")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--source", help="Path to SCF spreadsheet (.xlsx)")
    ap.add_argument("--db", help=f"Path to SQLite DB (default: {DEFAULT_DB})")
    ap.add_argument("--stats", action="store_true", help="Print stats and exit")
    args = ap.parse_args()

    db_path = resolve_db(args.db)
    if not db_path.exists():
        sys.exit(
            f"error: DB does not exist at {db_path}.\n"
            "  Run scripts/preprocess_scf.py first to build the base catalog."
        )

    if args.stats:
        with sqlite3.connect(db_path) as conn:
            print_stats(conn)
        return 0

    source = resolve_source(args.source)
    if not source.exists():
        sys.exit(f"error: source spreadsheet not found at {source}")

    # Defer import so the JSON-only preprocess path doesn't need openpyxl.
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("error: openpyxl is required. Install with `pip install openpyxl`.")

    print(f"Source: {source}")
    print(f"DB:     {db_path}")
    t0 = time.time()
    print("Loading workbook (read-only)…", flush=True)
    wb = load_workbook(source, read_only=True, data_only=True)
    print(f"  loaded in {time.time()-t0:.1f}s")

    print("Reading Authoritative Sources…", flush=True)
    authoritative = load_authoritative_sources(wb)
    print(f"  {len(authoritative)} framework entries indexed")

    print("Scanning SCF 2026.1 header row…", flush=True)
    ws_scf = wb["SCF 2026.1"]
    rows_iter = ws_scf.iter_rows(values_only=True)
    hdr = next(rows_iter)
    ctrl_id_col, framework_cols, cmm_cols = classify_headers(hdr, authoritative)
    if ctrl_id_col is None:
        sys.exit("error: could not locate 'SCF #' column in SCF 2026.1 sheet")
    print(
        f"  ctrl_id_col={ctrl_id_col + 1}  framework_cols={len(framework_cols)}  "
        f"cmm_cols={len(cmm_cols)}"
    )

    print("Augmenting DB…", flush=True)
    conn = sqlite3.connect(db_path)
    try:
        ensure_schema(conn)
        counts = populate(
            conn, authoritative, framework_cols, cmm_cols, ctrl_id_col, rows_iter
        )
        conn.commit()
    finally:
        conn.close()
    print(f"  walked {counts['rows_walked']} rows ({counts['rows_with_id']} with SCF #)")
    print(f"  framework_metadata inserted: {counts['framework_metadata']}")
    print(f"  control_frameworks inserted: {counts['control_frameworks']}")
    print(f"  maturity_criteria  inserted: {counts['maturity_criteria']}")
    print(f"  total time: {time.time() - t0:.1f}s")

    with sqlite3.connect(db_path) as conn:
        print()
        print_stats(conn)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
